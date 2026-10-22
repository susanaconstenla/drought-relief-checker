#!/usr/bin/env python3
"""
Combine drought relief disbursement files uploaded to the Cornell Box folder.

Reads every .xlsx and .csv file in a folder (for example your Box Drive folder, or a
folder you downloaded from Box), re-runs the same checks as the checker page, and
writes one combined dataset plus summaries.

Usage
-----
    python combine.py "/path/to/Box/Drought relief submissions"
    python combine.py "/path/to/folder" --out results --latest-per-org

Outputs (in --out, default ./combined_output)
--------------------------------------------
    combined.xlsx                 all sheets below, in one workbook
    rows_all.csv                  every row from every file, with status and issues
    rows_clean.csv                rows with no errors (use this for analysis)
    files_summary.csv             one line per file: who, counts, coverage
    coverage_ward_month.csv       125 wards x months: "aid", "zero", or blank (no record)
    run_log.txt                   what was read, skipped, and why

Rules match checker page version 2026-10-22. Requires: openpyxl (pip install openpyxl).
"""
import argparse
import csv
import json
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from datetime import date, datetime, timezone
from pathlib import Path

try:
    from openpyxl import Workbook, load_workbook
except ImportError:
    sys.exit("This script needs openpyxl. Install it with:  pip install openpyxl")

RULES_VERSION = "2026-10-22"

WARDS = ["Abakaile", "Adu", "Akithi", "Bahari", "Bamba", "Bangale", "Barwago", "Bulla Pesa", "Burat", "Bute", "Butiye", "Chala", "Chari", "Chepareria", "Cherab", "Chewani", "Churo/Amaya", "Dalalekutuk", "Dandu", "Dela", "Evurore", "Ewuaso O o Nkidong'i/M", "Faza", "Garsen Central", "Golbo", "Hadado/Adhibohol", "Iftin", "Ijara", "Ikutha", "Ilchamus", "Ilkerin", "Iloodokilani", "Imbirikani/Eselenkei", "Jaribuni", "Kaeris", "Kakuma", "Kakuyuni", "Kalaliyo", "Kangeta", "Kanyangi", "Kapchoki", "Kapedo/Napeitom", "Kapenguria", "Kaputei North", "Karare", "Kasemani", "Kasigau", "Kasikeu/Kiou", "Katilu", "Kauwi", "Kiambere", "Kilgoris Central", "Kinna", "Kipini East", "Kithungo/Kitundu", "Kitise/Kithuki", "Kiunga", "Kivou", "Kolowa", "Korr/Ngurunit", "Kutulo", "Lagboghol South", "Lake Zone", "Liboi", "Loiyangalani", "Lokichar", "Maalamin", "Madogo", "Magadi", "Maji Moto/Naroosura", "Malalani/Endau", "Masol", "Masongaleni", "Masumba/Nguu", "Mboghoni", "Merigi", "Migwani", "Mikinduri", "Mosiro (Narok East)", "Mugunda", "Mukaa", "Mukogodo East", "Mukogodo West", "Mukutani", "Mwanamwinga", "Mwarakaya", "Mwawesa", "Mwereni", "Nachola", "Nanighi", "Ndavaya", "Ndoto", "Ngobit", "Ngolia", "Nkareta", "Nuu", "Oldo-nyiro", "Rhamu Dimtu", "Ribikwo", "Riwo", "Rombo", "Sagante/Jaldessa", "Saimo/Kipsaraman", "Saka", "Sericho", "Shimbir Fatuma", "Siana", "Sosiani", "Suguta Marmar", "Tangulbei/Korossi", "Tarbaj", "Thegu River", "Tseikuru", "Turbi", "Turkwel", "Ukia", "Uran", "Waa", "Wamba North", "Wamba West", "Waranqara", "Waso", "Wayu", "Wei Wei", "Witu"]
COUNTIES = ["Baringo", "Embu", "Garissa", "Isiolo", "Kajiado", "Kilifi", "Kitui", "Kwale", "Laikipia", "Lamu", "Makueni", "Mandera", "Marsabit", "Meru", "Narok", "Nyeri", "Samburu", "Taita Taveta", "Tana River", "Turkana", "Wajir", "West Pokot"]
WARD_COUNTY = {"Churo/Amaya": "Baringo", "Ilchamus": "Baringo", "Kolowa": "Baringo", "Mukutani": "Baringo", "Ribikwo": "Baringo", "Saimo/Kipsaraman": "Baringo", "Tangulbei/Korossi": "Baringo", "Evurore": "Embu", "Kiambere": "Embu", "Abakaile": "Garissa", "Iftin": "Garissa", "Ijara": "Garissa", "Liboi": "Garissa", "Maalamin": "Garissa", "Nanighi": "Garissa", "Saka": "Garissa", "Bulla Pesa": "Isiolo", "Burat": "Isiolo", "Chari": "Isiolo", "Cherab": "Isiolo", "Kinna": "Isiolo", "Oldo-nyiro": "Isiolo", "Sericho": "Isiolo", "Dalalekutuk": "Kajiado", "Ewuaso O o Nkidong'i/M": "Kajiado", "Iloodokilani": "Kajiado", "Imbirikani/Eselenkei": "Kajiado", "Kaputei North": "Kajiado", "Magadi": "Kajiado", "Rombo": "Kajiado", "Adu": "Kilifi", "Bamba": "Kilifi", "Jaribuni": "Kilifi", "Kakuyuni": "Kilifi", "Mwanamwinga": "Kilifi", "Mwarakaya": "Kilifi", "Mwawesa": "Kilifi", "Ikutha": "Kitui", "Kanyangi": "Kitui", "Kauwi": "Kitui", "Kivou": "Kitui", "Malalani/Endau": "Kitui", "Migwani": "Kitui", "Nuu": "Kitui", "Tseikuru": "Kitui", "Kasemani": "Kwale", "Mwereni": "Kwale", "Ndavaya": "Kwale", "Waa": "Kwale", "Mukogodo East": "Laikipia", "Mukogodo West": "Laikipia", "Ngobit": "Laikipia", "Sosiani": "Laikipia", "Bahari": "Lamu", "Faza": "Lamu", "Kiunga": "Lamu", "Witu": "Lamu", "Kasikeu/Kiou": "Makueni", "Kithungo/Kitundu": "Makueni", "Kitise/Kithuki": "Makueni", "Masongaleni": "Makueni", "Masumba/Nguu": "Makueni", "Mukaa": "Makueni", "Ukia": "Makueni", "Dandu": "Mandera", "Kalaliyo": "Mandera", "Kutulo": "Mandera", "Rhamu Dimtu": "Mandera", "Shimbir Fatuma": "Mandera", "Waranqara": "Mandera", "Butiye": "Marsabit", "Golbo": "Marsabit", "Karare": "Marsabit", "Korr/Ngurunit": "Marsabit", "Loiyangalani": "Marsabit", "Merigi": "Marsabit", "Sagante/Jaldessa": "Marsabit", "Turbi": "Marsabit", "Uran": "Marsabit", "Akithi": "Meru", "Kangeta": "Meru", "Mikinduri": "Meru", "Ilkerin": "Narok", "Kilgoris Central": "Narok", "Maji Moto/Naroosura": "Narok", "Mosiro (Narok East)": "Narok", "Nkareta": "Narok", "Siana": "Narok", "Mugunda": "Nyeri", "Thegu River": "Nyeri", "Nachola": "Samburu", "Ndoto": "Samburu", "Suguta Marmar": "Samburu", "Wamba North": "Samburu", "Wamba West": "Samburu", "Waso": "Samburu", "Chala": "Taita Taveta", "Kasigau": "Taita Taveta", "Mboghoni": "Taita Taveta", "Ngolia": "Taita Taveta", "Bangale": "Tana River", "Chewani": "Tana River", "Garsen Central": "Tana River", "Kipini East": "Tana River", "Madogo": "Tana River", "Wayu": "Tana River", "Kaeris": "Turkana", "Kakuma": "Turkana", "Kapedo/Napeitom": "Turkana", "Katilu": "Turkana", "Lake Zone": "Turkana", "Lokichar": "Turkana", "Turkwel": "Turkana", "Barwago": "Wajir", "Bute": "Wajir", "Dela": "Wajir", "Hadado/Adhibohol": "Wajir", "Lagboghol South": "Wajir", "Tarbaj": "Wajir", "Chepareria": "West Pokot", "Kapchoki": "West Pokot", "Kapenguria": "West Pokot", "Masol": "West Pokot", "Riwo": "West Pokot", "Wei Wei": "West Pokot"}

# (key, column label in the template, output column name)
COLS = [
    ("county", "County", "county"),
    ("ward", "Ward", "ward"),
    ("month", "Month (YYYY-MM)", "month"),
    ("aid", "Type of aid distributed", "aid_type"),
    ("amount", "Amount disbursed", "amount"),
    ("unit", "Currency / unit", "amount_unit"),
    ("hh", "Total households served", "households"),
    ("agency", "Distributing agency", "agency"),
    ("pay", "Form of payment", "payment_form"),
    ("notes", "Notes", "notes"),
]
KEYS = [c[0] for c in COLS]
LABEL = {c[0]: c[1] for c in COLS}
OUTNAME = {c[0]: c[2] for c in COLS}

TODAY = date.today()
CUR = f"{TODAY.year}-{TODAY.month:02d}"
MONTHS = [f"{2024 + i // 12}-{i % 12 + 1:02d}" for i in range((TODAY.year - 2024) * 12 + TODAY.month)]
MNAMES = ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"]
NA = {"notapplicable", "nonapplicable", "noapplicable", "na", "notapp", "inkind"}
UNIT_NA = "Not applicable (in kind)"
UNIT_KES_WORDS = {"kes", "ksh", "kshs", "ks", "kenyashilling", "kenyashillings", "kenyanshilling", "kenyanshillings",
                  "shilling", "shillings"}
UNIT_USD_WORDS = {"usd", "us", "usdollar", "usdollars", "dollar", "dollars", "unitedstatesdollar", "unitedstatesdollars"}


def norm_unit(raw):
    """KES, USD, or 'Not applicable (in kind)'; '' if blank; None if not recognized (same as the page)."""
    t = str(raw if raw is not None else "").strip()
    if not t:
        return ""
    if t == "$":
        return "USD"
    n = norm(t)
    if n in UNIT_KES_WORDS:
        return "KES"
    if n in UNIT_USD_WORDS:
        return "USD"
    if n in NA or n in ("notapplicableinkind", "inkindnotapplicable"):
        return UNIT_NA
    return None


def norm(s):
    s = unicodedata.normalize("NFKD", str(s if s is not None else "").lower())
    return re.sub(r"[^a-z0-9]", "", s)


WARD_MAP = {norm(w): w for w in WARDS}
COUNTY_MAP = {norm(c): c for c in COUNTIES}


# ------------------------------------------------------------------ parsing (same as the page)
def excel_serial_to_month(n):
    d = date(1899, 12, 30).toordinal() + int(n)
    d = date.fromordinal(d)
    return f"{d.year}-{d.month:02d}"


def parse_month(raw):
    if isinstance(raw, (datetime, date)):
        return f"{raw.year}-{raw.month:02d}"
    s = str(raw if raw is not None else "").strip()
    if not s:
        return None
    if re.fullmatch(r"\d{5}(\.\d+)?", s):
        n = float(s)
        return excel_serial_to_month(n) if 30000 < n < 80000 else None
    m = None
    y = mo = None
    if (m := re.fullmatch(r"(\d{4})[-/. ](\d{1,2})", s)):
        y, mo = int(m[1]), int(m[2])
    elif (m := re.fullmatch(r"(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})(?:[ T].*)?", s)):
        y, mo = int(m[1]), int(m[2])
    elif (m := re.fullmatch(r"(\d{1,2})[-/.](\d{4})", s)):
        y, mo = int(m[2]), int(m[1])
    elif (m := re.fullmatch(r"(\d{4})(\d{2})", s)):
        y, mo = int(m[1]), int(m[2])
    elif (m := re.fullmatch(r"([A-Za-z]{3,9})\.?[\s,\-/]+(\d{4})", s)):
        y, mo = int(m[2]), (MNAMES.index(m[1][:3].lower()) + 1 if m[1][:3].lower() in MNAMES else 0)
    elif (m := re.fullmatch(r"(\d{4})[\s,\-/]+([A-Za-z]{3,9})", s)):
        y, mo = int(m[1]), (MNAMES.index(m[2][:3].lower()) + 1 if m[2][:3].lower() in MNAMES else 0)
    else:
        return None
    if not (1 <= mo <= 12) or y < 1990 or y > 2100:
        return None
    return f"{y}-{mo:02d}"


def num_str(x):
    """Number -> string the way the page shows it (3000.0 -> '3000', 3000.5 -> '3000.5')."""
    return str(int(x)) if float(x) == int(x) else repr(float(x))


def parse_amount(raw):
    s = str(raw if raw is not None else "").strip()
    if norm(s) in NA:
        return "not applicable"
    c = re.sub(r"(ksh|kes|usd|us\$|\$)", "", s, flags=re.I)
    c = re.sub(r"[,\s]", "", c)
    if re.fullmatch(r"-?\d+(\.\d+)?", c) or re.fullmatch(r"-?\.\d+", c):
        return float(c)
    return None


def parse_count(raw):
    c = re.sub(r"[,\s]", "", str(raw if raw is not None else "").strip())
    return int(float(c)) if re.fullmatch(r"\d+(\.0+)?", c) else None


def cell_text(v):
    if v is None:
        return ""
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, (int, float)):
        return num_str(v)
    if isinstance(v, (datetime, date)):
        return v.isoformat()
    return re.sub(r"\s+", " ", str(v)).strip()


def validate(v):
    """Normalize one row in place (dict of strings) and return (issues, fixes)."""
    iss, fixes = [], []
    add = lambda k, lvl, msg: iss.append((k, lvl, msg))
    if not v["county"]:
        add("county", "error", "County is missing")
    elif norm(v["county"]) in COUNTY_MAP:
        v["county"] = COUNTY_MAP[norm(v["county"])]
    else:
        add("county", "error", f"Not one of the {len(COUNTIES)} counties in scope")
    if not v["ward"]:
        add("ward", "error", "Ward is missing")
    elif norm(v["ward"]) in WARD_MAP:
        v["ward"] = WARD_MAP[norm(v["ward"])]
        if norm(v["county"]) in COUNTY_MAP and WARD_COUNTY[v["ward"]] != v["county"]:
            add("ward", "warning", f"{v['ward']} is in {WARD_COUNTY[v['ward']]} County, not {v['county']}")
    else:
        add("ward", "warning", f"Not one of the {len(WARDS)} wards in scope")
    if not v["month"]:
        add("month", "error", "Month is missing")
    else:
        m = parse_month(v["_month_raw"])
        if not m:
            add("month", "error", "Month not in YYYY-MM format")
        else:
            if m != v["month"]:
                fixes.append(("month", f'Month "{v["month"]}" read as {m}'))
                v["month"] = m
            if m > CUR:
                add("month", "error", "Month is in the future")
            elif m < "2024-01":
                add("month", "warning", "Before January 2024, outside the reporting window")
    no_aid = bool(re.match(r"no disbursement", v["aid"], re.I))
    if not v["aid"]:
        add("aid", "error", "Type of aid is missing")
    if v["amount"] == "":
        if no_aid:
            v["amount"] = "0"
            fixes.append(("amount", "Amount set to 0 for a month with no disbursement"))
        else:
            add("amount", "error", "Amount is missing")
    else:
        a = parse_amount(v["amount"])
        if a is None:
            add("amount", "error", "Amount is not a number or 'not applicable'")
        else:
            s = a if isinstance(a, str) else num_str(a)
            if s != v["amount"]:
                fixes.append(("amount", f'Amount "{v["amount"]}" read as {s}'))
                v["amount"] = s
            if not isinstance(a, str) and a < 0:
                add("amount", "error", "Amount is negative")
            if no_aid and a != 0:
                add("amount", "warning", "Type of aid says no disbursement, but the amount is not 0")
    if v["unit"]:
        u = norm_unit(v["unit"])
        if u is None:
            add("unit", "error", "Unit must be KES, USD, or Not applicable (in kind)")
        elif u != v["unit"]:
            fixes.append(("unit", f'Unit "{v["unit"]}" read as {u}'))
            v["unit"] = u
    zero = v["amount"] == "0"
    if not v["unit"] and not zero:
        add("unit", "error", "Currency / unit is missing")
    if v["unit"] == UNIT_NA and re.match(r"\d", v["amount"]) and not zero:
        add("amount", "warning", "Unit is in kind, so the amount should be 'not applicable'")
    if v["unit"] in ("KES", "USD") and v["amount"] == "not applicable":
        add("amount", "warning", f"Amount is 'not applicable' but the unit is {v['unit']}")
    if v["hh"] == "":
        if not zero:
            add("hh", "error", "Households served is missing")
    else:
        h = parse_count(v["hh"])
        if h is None:
            add("hh", "error", "Households is not a whole number")
        elif str(h) != v["hh"]:
            fixes.append(("hh", f'Households "{v["hh"]}" read as {h}'))
            v["hh"] = str(h)
    if len(v["notes"]) > 4000:
        add("notes", "error", "Notes longer than 4,000 characters")
    return iss, fixes


# ------------------------------------------------------------------ reading files
def col_key(h):
    n = norm(h)
    if not n:
        return None
    if n.startswith("ward"): return "ward"
    if n.startswith("county"): return "county"
    if n.startswith("month") or n in ("period", "date"): return "month"
    if "typeofaid" in n or n.startswith("aidtype") or n in ("aid", "typeofassistance"): return "aid"
    if n.startswith("amount"): return "amount"
    if n.startswith("currency") or n.startswith("unit"): return "unit"
    if "household" in n: return "hh"
    if any(t in n for t in ("agency", "implementing", "organization", "organisation")): return "agency"
    if "payment" in n or "schedule" in n: return "pay"
    if n.startswith("note") or n.startswith("comment"): return "notes"
    return None


def find_table(grid):
    """Return (header_row_index, {key: column_index}) or None."""
    for h, row in enumerate(grid[:15]):
        mp = {}
        for ci, cell in enumerate(row):
            k = col_key(cell)
            if k and k not in mp:
                mp[k] = ci
        if "ward" in mp and "month" in mp:
            return h, mp
    return None


def read_file(path):
    """Return dict with sheet, header row, rows [(excel_row, raw_cells)], mapping, info, or an error."""
    if path.suffix.lower() == ".csv":
        with open(path, newline="", encoding="utf-8-sig") as f:
            sheets = {"(csv)": list(csv.reader(f))}
        info = {}
    else:
        wb = load_workbook(path, read_only=True, data_only=True)
        sheets = {}
        names = sorted(wb.sheetnames, key=lambda n: norm(n) != "data")  # "Data" first
        for name in names:
            sheets[name] = [list(r) for r in wb[name].iter_rows(values_only=True)]
        info = {}
        for name in wb.sheetnames:
            if norm(name) == "submissioninfo":
                for r in sheets.get(name) or [list(x) for x in wb[name].iter_rows(values_only=True)]:
                    if r and r[0] not in (None, "Field"):
                        info[str(r[0]).strip()] = cell_text(r[1] if len(r) > 1 else "")
        wb.close()
    for name, grid in sheets.items():
        if norm(name) == "submissioninfo" or norm(name) == "checkernotes":
            continue
        t = find_table(grid)
        if not t:
            continue
        h, mp = t
        rows = []
        for i in range(h + 1, len(grid)):
            line = grid[i]
            raw = {k: (line[mp[k]] if k in mp and mp[k] < len(line) else None) for k in KEYS}
            if all(cell_text(raw[k]) == "" for k in KEYS):
                continue
            rows.append((i + 1, raw))  # Excel row number
        missing = [k for k in KEYS if k not in mp]
        return {"sheet": name, "rows": rows, "missing": missing, "info": info}
    return {"error": "No table with 'Ward' and 'Month' headings found"}


# ------------------------------------------------------------------ main
def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("folder", help="Folder with the uploaded files (searched recursively)")
    ap.add_argument("--out", default="combined_output", help="Output folder (default: combined_output)")
    ap.add_argument("--latest-per-org", action="store_true",
                    help="Keep only the most recent file per organization (use when partners re-upload corrections)")
    a = ap.parse_args()

    folder = Path(a.folder).expanduser()
    if not folder.is_dir():
        sys.exit(f"Folder not found: {folder}")
    out = Path(a.out)
    out.mkdir(parents=True, exist_ok=True)
    log = [f"Run {datetime.now(timezone.utc).isoformat(timespec='seconds')}  rules {RULES_VERSION}  folder {folder}"]

    files = sorted(p for p in folder.rglob("*")
                   if p.is_file() and p.suffix.lower() in (".xlsx", ".csv") and not p.name.startswith(("~$", ".")))
    skipped = sorted(p for p in folder.rglob("*") if p.is_file() and p.suffix.lower() in (".xls", ".xlsm", ".ods"))
    for p in skipped:
        log.append(f"SKIPPED {p.name}: save it as .xlsx and run again")

    parsed = []
    for p in files:
        try:
            r = read_file(p)
        except Exception as e:  # unreadable file: log and keep going
            log.append(f"SKIPPED {p.name}: could not read ({e})")
            continue
        if "error" in r:
            log.append(f"SKIPPED {p.name}: {r['error']}")
            continue
        info = r["info"]
        org = info.get("Organization") or p.stem
        checked = info.get("Checked on (UTC)", "")
        stamp = checked or datetime.fromtimestamp(p.stat().st_mtime, timezone.utc).isoformat(timespec="seconds")
        parsed.append({"path": p, "org": org, "org_key": norm(org), "stamp": stamp, **r})
        log.append(f"READ {p.name}: sheet '{r['sheet']}', {len(r['rows'])} rows, organization '{org}'"
                   + (f", missing columns: {', '.join(LABEL[k] for k in r['missing'])}" if r["missing"] else ""))

    if a.latest_per_org:
        latest = {}
        for f in parsed:
            if f["org_key"] not in latest or f["stamp"] > latest[f["org_key"]]["stamp"]:
                latest[f["org_key"]] = f
        dropped = [f for f in parsed if latest[f["org_key"]] is not f]
        for f in dropped:
            log.append(f"NOT USED {f['path'].name}: a newer file from '{f['org']}' exists (--latest-per-org)")
        parsed = [f for f in parsed if latest[f["org_key"]] is f]

    # validate
    records = []
    for f in parsed:
        for excel_row, raw in f["rows"]:
            v = {k: cell_text(raw[k]) for k in KEYS}
            v["_month_raw"] = raw["month"] if raw["month"] is not None else ""
            iss, fixes = validate(v)
            records.append({"file": f["path"].name, "org": f["org"], "excel_row": excel_row, "v": v,
                            "issues": iss, "fixes": fixes})

    # duplicates: within a file, and across files
    def dkey(v):
        return "|".join(norm(v[k]) for k in ("ward", "month", "aid", "agency"))
    by_key = defaultdict(list)
    for r in records:
        if r["v"]["ward"] and r["v"]["month"]:
            by_key[dkey(r["v"])].append(r)
    for group in by_key.values():
        if len(group) < 2:
            continue
        files_in = {r["file"] for r in group}
        for r in group:
            if sum(1 for x in group if x["file"] == r["file"]) > 1:
                r["issues"].append(("_dup", "warning", "Same ward, month, aid type and agency as another row in this file"))
            if len(files_in) > 1:
                others = sorted(files_in - {r["file"]})
                r["issues"].append(("_dup", "warning", "Same ward, month, aid type and agency also in: " + "; ".join(others)))

    def status(r):
        lv = {i[1] for i in r["issues"]}
        return "error" if "error" in lv else ("warning" if "warning" in lv else "ok")

    def issues_text(r):
        parts = [f"{i[1].upper()}: {i[2]}" for i in r["issues"]] + [f"FIXED: {m}" for _, m in r["fixes"]]
        return " | ".join(parts)

    head = ["source_file", "organization", "file_row"] + [OUTNAME[k] for k in KEYS] + ["status", "issues"]

    def as_row(r):
        v = r["v"]
        return [r["file"], r["org"], r["excel_row"]] + [v[k] for k in KEYS] + [status(r), issues_text(r)]

    all_rows = [as_row(r) for r in records]
    clean = [r for r in records if status(r) != "error"]
    clean_rows = [as_row(r) for r in clean]

    # per-file summary
    summ_head = ["source_file", "organization", "contact_person", "contact_email", "checked_on_page", "rows",
                 "rows_ok", "rows_warning", "rows_error", "wards_covered", "months_covered", "first_month",
                 "last_month", "households_total", "message"]
    summ = []
    for f in parsed:
        rs = [r for r in records if r["file"] == f["path"].name]
        ok = [r for r in rs if status(r) != "error"]
        months = sorted({r["v"]["month"] for r in ok if re.fullmatch(r"\d{4}-\d{2}", r["v"]["month"])})
        hh = sum(parse_count(r["v"]["hh"]) or 0 for r in ok)
        info = f["info"]
        summ.append([f["path"].name, f["org"], info.get("Contact person", ""), info.get("Contact email", ""),
                     "yes" if info.get("Checker version") else "no", len(rs),
                     sum(status(r) == "ok" for r in rs), sum(status(r) == "warning" for r in rs),
                     sum(status(r) == "error" for r in rs),
                     len({r["v"]["ward"] for r in ok if r["v"]["ward"] in WARD_MAP.values()}), len(months),
                     months[0] if months else "", months[-1] if months else "", hh,
                     info.get("Message to the research team", "")])

    # coverage: ward x month from rows without errors; "aid" beats "zero"
    cov = defaultdict(dict)
    for r in clean:
        w, m = r["v"]["ward"], r["v"]["month"]
        if w in WARD_MAP.values() and m in MONTHS:
            val = "zero" if r["v"]["amount"] == "0" else "aid"
            if cov[w].get(m) != "aid":
                cov[w][m] = val
    cov_head = ["ward"] + MONTHS
    cov_rows = [[w] + [cov[w].get(m, "") for m in MONTHS] for w in WARDS]

    # write CSVs
    def write_csv(name, header, rows):
        with open(out / name, "w", newline="", encoding="utf-8") as fh:
            wr = csv.writer(fh)
            wr.writerow(header)
            wr.writerows(rows)
    write_csv("rows_all.csv", head, all_rows)
    write_csv("rows_clean.csv", head, clean_rows)
    write_csv("files_summary.csv", summ_head, summ)
    write_csv("coverage_ward_month.csv", cov_head, cov_rows)

    # write workbook (text that looks like a formula stays text)
    wb = Workbook()
    first = True
    for title, header, rows in [("Rows (clean)", head, clean_rows), ("Rows (all)", head, all_rows),
                                ("Files", summ_head, summ), ("Coverage", cov_head, cov_rows)]:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = title
        ws.append(header)
        for row in rows:
            ws.append(row)
            for cell in ws[ws.max_row]:
                if isinstance(cell.value, str) and cell.value.startswith(("=", "+", "-", "@")):
                    cell.data_type = "s"
        ws.freeze_panes = "A2"
    wb.save(out / "combined.xlsx")

    n_err = sum(status(r) == "error" for r in records)
    n_warn = sum(status(r) == "warning" for r in records)
    wards_cov = sum(1 for w in WARDS if cov[w])
    log += ["", f"Files used: {len(parsed)}", f"Rows read: {len(records)}",
            f"Rows with errors (left out of rows_clean.csv): {n_err}",
            f"Rows with warnings (kept, check 'issues'): {n_warn}",
            f"Wards with at least one month: {wards_cov} of {len(WARDS)}",
            f"Output folder: {out.resolve()}"]
    (out / "run_log.txt").write_text("\n".join(log) + "\n", encoding="utf-8")
    print("\n".join(log))


if __name__ == "__main__":
    main()
